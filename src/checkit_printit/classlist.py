"""Read a class list in whatever shape the registrar or the LMS produced.

Three real exports drove this, and no two agree:

| | Banner detail CSV | LMS classlist CSV | Banner summary XLSX |
|---|---|---|---|
| header row | first | first | **fifteenth**, under a preamble |
| name | three columns | `Last, First` in one | `Last, First M.` in one |
| id | `Student ID` `806...` | `OrgDefinedId` `20...` | `ID` `806...` |
| second id | `Global ID` `20...` | -- | -- |
| email | yes | yes | **none** |
| section | yes | no | no |
| non-students | -- | **an instructor and a mentor row** | `Registration Status` |

So: find the header row rather than assume it, match columns by label against
a synonym table, drop rows that are not students, and say what was decided
instead of deciding silently.

**Both id systems are kept.** They are different numbers for the same person,
and an export carries one or the other, so a merge matches on whichever it
has. Email is a weak key and not a strong one: the same student appears as
`mcliffo4@oswego.edu` in one of these files and `m.clifford@clasnet...` in
another, downloaded the same day.
"""

import csv
import dataclasses
import os
import re

from .roster import Roster, RosterError, Student

#: Column labels seen in the wild, lower-cased and stripped. Order inside a
#: list is preference: the first match wins when a file offers several.
SYNONYMS = {
    "sid": ["student id", "sid", "banner id", "id"],
    "alt_id": ["global id", "orgdefinedid", "org defined id", "user id",
               "username", "user name"],
    "last": ["student last name", "last name", "last", "surname",
             "family name"],
    "first": ["student first name", "first name", "first", "given name"],
    "middle": ["student mi", "mi", "middle initial", "middle name", "middle"],
    "full": ["student name", "full name", "name"],
    "email": ["email", "email address", "e-mail", "student email",
              "primary email"],
    "section": ["section", "sec", "section number", "course section"],
    "role": ["role", "registration status", "status", "enrollment status",
             "enrolment status"],
}

#: A role or status that means this row is not a student to print for.
NOT_A_STUDENT = ("instructor", "teacher", "professor", "mentor", "assistant",
                 "designer", "observer", "auditor", "grader", "staff",
                 "dropped", "withdrawn", "deleted", "waitlist", "wait list")

#: A role or status that confirms one is.
IS_A_STUDENT = ("student", "registered", "enrolled", "active", "learner")


class ClassListError(Exception):
    pass


@dataclasses.dataclass
class Mapping:
    """Which column held what, and how sure we are."""
    header_row: int
    columns: dict            # field -> column index
    unmatched: list          # header labels we did not recognise

    def describe(self):
        lines = [f"header row {self.header_row + 1}"]
        for field in ("sid", "alt_id", "last", "first", "middle", "full",
                      "email", "section", "role"):
            if field in self.columns:
                lines.append(f"  {field:8} column {self.columns[field] + 1}")
        if self.unmatched:
            lines.append("  ignored: " + ", ".join(self.unmatched))
        return "\n".join(lines)


# ---------------------------------------------------------------- reading --

def read_rows(path):
    """Every cell of the file as strings, whatever the format."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".tsv", ".txt"):
        delimiter = "\t" if ext == ".tsv" else ","
        with open(path, encoding="utf-8-sig", newline="") as f:
            return [[(c or "").strip() for c in row]
                    for row in csv.reader(f, delimiter=delimiter)]
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:                                  # pragma: no cover
            raise ClassListError(
                "reading .xlsx needs openpyxl, which is not installed. "
                "Install it, or save the file as CSV first."
            ) from None
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            book = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet = book[book.sheetnames[0]]
        rows = [["" if c.value is None else str(c.value).strip() for c in row]
                for row in sheet.iter_rows()]
        book.close()
        return rows
    raise ClassListError(
        f"{path}: unrecognised extension {ext!r}. Expected .csv, .tsv or .xlsx."
    )


def _label(cell):
    return re.sub(r"\s+", " ", cell.strip().lower()).strip(" :*")


def find_mapping(rows):
    """Locate the header row and map its columns onto known fields.

    The header is not always the first row -- a Banner summary puts a course
    information block above it -- so every row is scored and the best wins.
    """
    best = None
    for r, row in enumerate(rows[:60]):
        labels = [_label(c) for c in row]
        columns, used = {}, set()
        for field, options in SYNONYMS.items():
            for option in options:
                if option in labels:
                    index = labels.index(option)
                    if index in used:
                        continue
                    columns[field] = index
                    used.add(index)
                    break
        # a header row has to name somebody and be more than one column
        nameable = ("full" in columns) or ("last" in columns)
        if not nameable or len(columns) < 2:
            continue
        unmatched = [c for i, c in enumerate(row)
                     if i not in used and c.strip()]
        candidate = Mapping(header_row=r, columns=columns, unmatched=unmatched)
        if best is None or len(columns) > len(best.columns):
            best = candidate
    if best is None:
        raise ClassListError(
            "could not find a header row. Looked for a row naming a student "
            "-- 'Student Name', 'Last Name' or similar -- alongside at least "
            "one of an id, an email or a section."
        )
    return best


# ---------------------------------------------------------------- parsing --

def split_full_name(value):
    """`Last, First M.` -> (last, first, middle). Also handles `First Last`.

    The comma form is what both single-column exports use, and it is the one
    that gets multi-word surnames right: "Soriano Garcia, Belinda" has to keep
    both words on the left.
    """
    value = re.sub(r"\s+", " ", value).strip()
    if not value:
        return "", "", ""
    if "," in value:
        last, _, rest = value.partition(",")
        parts = rest.strip().split()
        first = parts[0] if parts else ""
        middle = " ".join(parts[1:])
        return last.strip(), first, middle
    parts = value.split()
    if len(parts) == 1:
        return parts[0], "", ""
    return parts[-1], parts[0], " ".join(parts[1:-1])


def looks_like_a_student(value):
    """(keep, reason). An unrecognised role is kept, and reported."""
    text = _label(value)
    if not text:
        return True, ""
    for bad in NOT_A_STUDENT:
        if bad in text:
            return False, value.strip()
    for good in IS_A_STUDENT:
        if good in text:
            return True, ""
    return True, f"unrecognised role {value.strip()!r}"


def parse(path, section=""):
    """Read a class list. Returns (students, mapping, notes)."""
    rows = read_rows(path)
    if not rows:
        raise ClassListError(f"{path} is empty.")
    mapping = find_mapping(rows)
    col = mapping.columns
    notes = []

    def cell(row, field):
        i = col.get(field)
        if i is None or i >= len(row):
            return ""
        return row[i].strip()

    students, skipped = [], []
    for row in rows[mapping.header_row + 1:]:
        if not any(c.strip() for c in row):
            continue

        keep, reason = looks_like_a_student(cell(row, "role"))
        if not keep:
            skipped.append(reason)
            continue
        if reason:
            notes.append(reason)

        last, first, middle = cell(row, "last"), cell(row, "first"), cell(row, "middle")
        if not (last or first):
            last, first, middle = split_full_name(cell(row, "full"))
        if not (last or first):
            continue                       # a spacer row, not a person

        students.append(Student(
            name=f"{first} {last}".strip(),
            skills=[],
            last=last,
            first=first,
            preferred=first,
            sid=cell(row, "sid"),
            alt_id=cell(row, "alt_id"),
            email=cell(row, "email").lower(),
            section=cell(row, "section") or section,
        ))

    if not students:
        raise ClassListError(
            f"{path}: found a header row but no students under it.\n"
            + mapping.describe()
        )

    if skipped:
        notes.append(f"skipped {len(skipped)} non-student row(s): "
                     + ", ".join(sorted(set(skipped))))
    if not any(s.email for s in students):
        notes.append(
            "no email column, so these students cannot be matched to form "
            "responses. Import a file that has one before pulling responses."
        )
    if not any(s.section for s in students):
        notes.append("no section column; pass --section to set one.")
    return students, mapping, notes


# ----------------------------------------------------------------- merging --

@dataclasses.dataclass
class MergeReport:
    added: list
    updated: list
    dropped: list
    unchanged: list
    kept_dropped: list = dataclasses.field(default_factory=list)

    def describe(self):
        out = (f"{len(self.added)} added, {len(self.updated)} updated, "
               f"{len(self.dropped)} marked dropped, "
               f"{len(self.unchanged)} unchanged")
        if self.kept_dropped:
            out += (f", {len(self.kept_dropped)} still dropped by hand "
                    f"despite being listed")
        return out


def _norm(value):
    return " ".join((value or "").split()).lower()


def _as_first_last(student):
    """(first, last), falling back to splitting the display name.

    A roster adopted from an old print job has only `name`, so without this
    the surname-plus-initial rule has nothing on its side of the comparison
    and every preferred name -- Matt for Matthew, Aly for Alyson -- imports as
    a second copy of a student already there.
    """
    if student.first and student.last:
        return student.first, student.last
    parts = (student.name or "").split(None, 1)
    if len(parts) == 2:
        return parts[0], parts[1]
    return student.first, student.last


def _index(students):
    """Every key to the student holding it, plus the ambiguous ones removed.

    Names are weak keys and are only reached when no id matched. Two are kept:
    the display name, because a roster adopted from an old print job has
    nothing else, and surname-plus-initial, because a chart says "Matt
    Brienza" where a class list says "Brienza, Matthew C.".

    Surname-plus-initial can collide -- two Smiths, both J. A colliding key is
    dropped rather than resolved, so a merge never guesses which one it meant.
    """
    out = {}
    seen_initials = {}
    for s in students:
        for kind, value in s.ids():
            if value:
                out[(kind, value)] = s
        for spelling in (s.name, f"{s.first} {s.last}", f"{s.preferred} {s.last}"):
            key = _norm(spelling)
            if key:
                out.setdefault(("name", key), s)
        first, last = _as_first_last(s)
        if first and last:
            initial = ("initial", _norm(last), first[:1].lower())
            seen_initials.setdefault(initial, []).append(s)
    for key, holders in seen_initials.items():
        if len(holders) == 1:
            out[key] = holders[0]
    return out


def _spellings(student):
    """The name keys to try for an incoming student, strongest first."""
    keys = []
    for spelling in (f"{student.first} {student.last}", student.name):
        key = _norm(spelling)
        if key:
            keys.append(("name", key))
    first, last = _as_first_last(student)
    if first and last:
        keys.append(("initial", _norm(last), first[:1].lower()))
    return keys


def merge(existing, incoming, scope=None):
    """Fold a freshly read class list into the roster we already have.

    Absence is not deletion: a student the import does not mention is marked
    dropped and kept, because the print record refers to them. And nothing the
    instructor authored is overwritten -- `preferred`, `name` and `skills`
    survive every re-import, which is what stops a class list renaming the
    person on the printed page.

    **Absence only counts inside the sections the file covers.** A class list
    is usually one section, and a course may hold several; without this,
    importing 820 would mark all of 830 as dropped. `scope` overrides the
    sections inferred from the file, and an import carrying no section at all
    is taken to cover everyone, because nothing says otherwise.
    """
    students = [dataclasses.replace(s) for s in existing]
    index = _index(students)
    report = MergeReport([], [], [], [], [])
    seen = set()

    for fresh in incoming:
        match = None
        for kind, value in fresh.ids():
            if value and (kind, value) in index:
                match = index[(kind, value)]
                break
        if match is None:
            for key in _spellings(fresh):
                if key in index:
                    match = index[key]
                    break

        if match is None:
            students.append(fresh)
            report.added.append(fresh.name)
            seen.add(id(students[-1]))
            continue

        seen.add(id(match))
        changed = []
        for field in ("last", "first", "sid", "alt_id", "section"):
            new = getattr(fresh, field)
            if new and new != getattr(match, field):
                changed.append(field)
                setattr(match, field, new)

        # An address is added, never substituted. Which one the Google Form
        # sees is not ours to guess, so every one a student has been seen
        # under stays a way to find them.
        for address in fresh.all_emails():
            if address not in match.all_emails():
                if not match.email:
                    match.email = address
                    changed.append("email")
                else:
                    match.emails = list(match.emails) + [address]
                    changed.append(f"second address {address}")
        if match.dropped and match.dropped_by != "instructor":
            match.dropped = False
            match.dropped_by = ""
            changed.append("re-enrolled")
        elif match.dropped:
            report.kept_dropped.append(match.name)
        if not match.preferred:
            match.preferred = match.first
        if changed:
            report.updated.append(f"{match.name} ({', '.join(changed)})")
        else:
            report.unchanged.append(match.name)

    if scope is None:
        scope = {s.section for s in incoming if s.section}
    for s in students:
        if id(s) in seen or s.dropped:
            continue
        if scope and s.section and s.section not in scope:
            continue                      # a section this file says nothing about
        s.dropped = True
        s.dropped_by = "import"
        report.dropped.append(s.name)

    return Roster(students), report
